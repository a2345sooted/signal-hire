import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

class DOCXGenerator:
    def __init__(self):
        pass

    def generate_docx(self, resume_data: dict) -> bytes:
        doc = Document()
        
        # Set margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # 1. Contact Info
        contact = resume_data.get("contact", {})
        if contact.get("name"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(contact["name"])
            run.bold = True
            run.font.size = Pt(18)
        
        contact_parts = []
        if contact.get("email"): contact_parts.append(contact["email"])
        if contact.get("phone"): contact_parts.append(contact["phone"])
        if contact.get("location"): contact_parts.append(contact["location"])
        if contact.get("linkedin"): contact_parts.append(contact["linkedin"])
        
        if contact_parts:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(" | ".join(contact_parts)).font.size = Pt(10)

        # 2. Summary
        if resume_data.get("summary"):
            self._add_section_heading(doc, "PROFESSIONAL SUMMARY")
            doc.add_paragraph(resume_data["summary"])

        # 3. Experience
        if resume_data.get("experience"):
            self._add_section_heading(doc, "PROFESSIONAL EXPERIENCE")
            for exp in resume_data["experience"]:
                p = doc.add_paragraph()
                title = exp.get('title', '')
                company = exp.get('company', '')
                
                run = p.add_run(title)
                run.bold = True
                if company:
                    p.add_run(f" | {company}")
                
                if exp.get('dates'):
                    p_dates = doc.add_paragraph()
                    run_dates = p_dates.add_run(exp['dates'])
                    run_dates.italic = True
                    run_dates.font.size = Pt(10)
                    p_dates.paragraph_format.space_after = Pt(2)

                bullets = exp.get("bullets", [])
                for bullet in bullets:
                    doc.add_paragraph(bullet, style='List Bullet')

        # 4. Projects
        projects = resume_data.get("projects", [])
        ai_projects = resume_data.get("ai_projects", [])
        all_projects = projects + ai_projects
        
        if all_projects:
            self._add_section_heading(doc, "PROJECTS")
            for proj in all_projects:
                p = doc.add_paragraph()
                run = p.add_run(proj.get('name', ''))
                run.bold = True
                if proj.get('technologies'):
                    p.add_run(f" ({', '.join(proj['technologies'])})")
                
                if proj.get('description'):
                    doc.add_paragraph(proj['description'])
                
                for link in proj.get('links', []):
                    doc.add_paragraph(f"Link: {link}")

        # 5. Skills
        if resume_data.get("skills"):
            self._add_section_heading(doc, "SKILLS")
            doc.add_paragraph(", ".join(resume_data["skills"]))

        # 6. Education
        if resume_data.get("education"):
            self._add_section_heading(doc, "EDUCATION")
            for edu in resume_data["education"]:
                p = doc.add_paragraph()
                run = p.add_run(edu.get('degree', ''))
                run.bold = True
                school = edu.get('school', '')
                dates = edu.get('dates', '')
                
                suffix = ""
                if school:
                    suffix += f", {school}"
                if dates:
                    suffix += f" ({dates})"
                
                if suffix:
                    p.add_run(suffix)

        # 7. Military Service
        if resume_data.get("military_service"):
            self._add_section_heading(doc, "MILITARY SERVICE")
            doc.add_paragraph(resume_data["military_service"])

        # 8. Interests
        if resume_data.get("interests"):
            self._add_section_heading(doc, "INTERESTS & CERTIFICATIONS")
            doc.add_paragraph(", ".join(resume_data["interests"]))

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes

    def _add_section_heading(self, doc, text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)

class DOCXGeneratorService(DOCXGenerator):
    pass
